"""Word文档渲染器"""
from __future__ import annotations
from typing import Optional

import io
import logging
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

logger = logging.getLogger(__name__)


def render_docx(content: dict, output_path: Path):
    """将文档JSON渲染为Word文件"""
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)

    title = content.get("title", "未命名文档")
    sections = content.get("sections", [])

    for section in sections:
        section_type = section.get("type")

        if section_type == "heading":
            level = section.get("level", 1)
            text = section.get("text", "")
            heading = doc.add_heading(text, level=level)
            # 设置标题字体
            for run in heading.runs:
                run.font.name = '黑体'

        elif section_type == "paragraph":
            text = section.get("text", "")
            doc.add_paragraph(text)

        elif section_type == "bullet":
            text = section.get("text", "")
            items = section.get("items", [])
            if items:
                for item in items:
                    if isinstance(item, str):
                        doc.add_paragraph(item, style='List Bullet')
                    else:
                        doc.add_paragraph(str(item), style='List Bullet')
            elif text:
                doc.add_paragraph(text, style='List Bullet')

        elif section_type == "table":
            headers = section.get("headers", [])
            rows = section.get("rows", [])
            if headers:
                table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                table.style = 'Light Grid Accent 1'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # 表头
                for i, header in enumerate(headers):
                    cell = table.rows[0].cells[i]
                    cell.text = str(header)
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in paragraph.runs:
                            run.bold = True

                # 数据行
                for row_idx, row_data in enumerate(rows):
                    for col_idx, cell_data in enumerate(row_data):
                        if col_idx < len(headers):
                            table.rows[row_idx + 1].cells[col_idx].text = str(cell_data)

        elif section_type == "chart":
            # 图表：用matplotlib生成图片后嵌入
            chart_data = section.get("data", {})
            chart_title = section.get("title", "图表")
            chart_type = section.get("chart_type", "bar")

            try:
                img_path = _render_chart_image(chart_data, chart_title, chart_type, output_path)
                if img_path:
                    doc.add_picture(str(img_path), width=Inches(5.5))
                    # 图片居中
                    last_paragraph = doc.paragraphs[-1]
                    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                logger.error(f"图表渲染失败: {e}")
                doc.add_paragraph(f"[图表渲染失败: {chart_title}]")

        elif section_type == "page_break":
            doc.add_page_break()

    # 保存
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    logger.info(f"Word文档已生成: {output_path}")


def _render_chart_image(chart_data: dict, title: str, chart_type: str,
                        doc_path: Path) -> Path | None:
    """用matplotlib渲染图表为图片"""
    import matplotlib
    matplotlib.use('Agg')  # 无GUI后端
    import matplotlib.pyplot as plt

    plt.rcParams['font.sans-serif'] = ['SimHei', 'DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False

    labels = chart_data.get("labels", [])
    datasets = chart_data.get("datasets", [])

    fig, ax = plt.subplots(figsize=(8, 5))

    if chart_type == "bar":
        x = range(len(labels))
        for ds in datasets:
            ax.bar(x, ds.get("values", []), label=ds.get("label", ""))
        ax.set_xticks(x)
        ax.set_xticklabels(labels)
    elif chart_type == "line":
        for ds in datasets:
            ax.plot(labels, ds.get("values", []), marker='o', label=ds.get("label", ""))
    elif chart_type == "pie":
        if datasets:
            values = datasets[0].get("values", [])
            ax.pie(values, labels=labels, autopct='%1.1f%%')
    elif chart_type == "scatter":
        if datasets:
            values = datasets[0].get("values", [])
            ax.scatter(labels[:len(values)], values)
    else:
        logger.warning(f"不支持的图表类型: {chart_type}")
        return None

    if chart_type != "pie":
        ax.legend()

    ax.set_title(title)
    fig.tight_layout()

    # 保存图片
    img_path = doc_path.parent / f"chart_{hash(title) & 0xFFFFFFFF}.png"
    fig.savefig(str(img_path), dpi=150, bbox_inches='tight')
    plt.close(fig)

    return img_path
