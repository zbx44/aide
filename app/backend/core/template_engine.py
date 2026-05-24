from __future__ import annotations
"""文档模板引擎 - 解析模板、填充变量、AI生成内容、渲染文档

模板系统架构：
  模板 = Word/PPT文件 + 变量定义(template.yaml) + AI Prompt

  模板来源：
  - config/templates/  ← 系统内置模板（升级时覆盖）
  - user/templates/    ← 用户自建模板（升级不碰）

  使用流程：
  1. 选择模板
  2. 填写变量（部分自动从数据库获取）
  3. AI根据Prompt填充文档内容
  4. 渲染为Word/Excel/PPT文件
"""
import json
import logging
import re
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Optional

import yaml

from config.settings import settings
from core.llm import llm_client

logger = logging.getLogger(__name__)


def _sections_to_slides(sections, title="未命名"):
    """从sections构建PPT幻灯片列表"""
    slides = [{"layout": "title", "title": title}]
    current_body = []
    current_title = ""

    for section in sections:
        if section.get("type") == "heading" and section.get("level", 1) <= 2:
            if current_body or current_title:
                slides.append({"layout": "content", "title": current_title, "body": current_body})
            current_title = section.get("text", "")
            current_body = []
        elif section.get("type") in ("paragraph", "bullet"):
            entry = {
                "type": "text" if section["type"] == "paragraph" else "bullet",
                "content": section.get("text", ""),
            }
            if section["type"] == "bullet":
                entry["items"] = section.get("items", [section.get("text", "")])
            current_body.append(entry)
        elif section.get("type") == "table":
            current_body.append({"type": "text", "content": json.dumps(section, ensure_ascii=False)[:200]})
        elif section.get("type") == "chart":
            if current_body or current_title:
                slides.append({"layout": "content", "title": current_title, "body": current_body})
                current_body = []
            slides.append({"layout": "chart", "title": section.get("title", "图表"), "chart": section})
            current_title = ""

    if current_body or current_title:
        slides.append({"layout": "content", "title": current_title, "body": current_body})

    return slides


def _slides_to_sections(slides):
    """从slides反向构建sections列表"""
    sections = []
    for slide in slides:
        if slide.get("layout") == "title":
            sections.append({"type": "heading", "level": 1, "text": slide.get("title", "")})
            if slide.get("subtitle"):
                sections.append({"type": "paragraph", "text": slide["subtitle"]})
        else:
            sections.append({"type": "heading", "level": 2, "text": slide.get("title", "")})
            for item in slide.get("body", []):
                if item.get("type") == "bullet":
                    for b in item.get("items", []):
                        sections.append({"type": "bullet", "text": b})
                else:
                    sections.append({"type": "paragraph", "text": item.get("content", "")})
    return sections


class TemplateEngine:
    """文档模板引擎"""

    def __init__(self):
        self.system_templates_dir = settings.CONFIG_DIR / "templates"
        self.user_templates_dir = settings.USER_DIR / "templates"
        self.output_dir = settings.OUTPUT_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def list_templates(self, category: str = None) -> List[dict]:
        """列出所有可用模板（系统+用户）"""
        templates = []
        seen_names = set()

        # 用户模板优先
        for tmpl_dir in self._scan_template_dirs(self.user_templates_dir):
            tmpl = self._load_template_info(tmpl_dir, scope="personal")
            if tmpl and tmpl["name"] not in seen_names:
                templates.append(tmpl)
                seen_names.add(tmpl["name"])

        # 系统模板
        for tmpl_dir in self._scan_template_dirs(self.system_templates_dir):
            tmpl = self._load_template_info(tmpl_dir, scope="system")
            if tmpl and tmpl["name"] not in seen_names:
                templates.append(tmpl)
                seen_names.add(tmpl["name"])

        # 按分类过滤
        if category and category != "全部":
            templates = [t for t in templates if t.get("category") == category]

        return templates

    def get_template(self, name: str) -> Optional[dict]:
        """获取模板详情"""
        # 先查用户模板，再查系统模板
        for base_dir in [self.user_templates_dir, self.system_templates_dir]:
            tmpl_dir = base_dir / name
            if tmpl_dir.exists():
                return self._load_template_info(tmpl_dir)
        return None

    def generate_document(self, template_name: str, variables: dict) -> dict:
        """根据模板生成文档

        Args:
            template_name: 模板名称
            variables: 用户填写的变量值

        Returns:
            {"doc_id": "xxx", "file_path": "...", "title": "...", "status": "success"}
            或 {"error": "...", "status": "error"}
        """
        # 1. 加载模板
        template = self._load_full_template(template_name)
        if not template:
            return {"error": f"模板 '{template_name}' 不存在", "status": "error"}

        # 2. 校验必填变量
        missing = []
        for var_def in template.get("variables", []):
            if var_def.get("required") and not variables.get(var_def["name"]):
                missing.append(var_def.get("label", var_def["name"]))
        if missing:
            return {"error": f"请填写必填项：{', '.join(missing)}", "status": "error"}

        # 3. 处理变量（自动填充数据库查询类型的变量）
        resolved_vars = self._resolve_variables(template, variables)

        # 4. AI填充内容
        try:
            ai_content = self._ai_fill(template, resolved_vars)
        except Exception as e:
            logger.error(f"AI填充失败: {e}", exc_info=True)
            return {"error": f"AI生成内容失败: {str(e)[:200]}", "status": "error"}

        # 5. 渲染文档
        doc_id = str(uuid.uuid4())[:8]
        doc_type = template.get("doc_type", "docx")

        # 生成文件名
        filename_template = template.get("output", {}).get(
            "filename", f"{template_name}_{{{{date}}}}.{doc_type}"
        )
        filename = self._render_filename(filename_template, resolved_vars, doc_type)

        # 保存目录
        save_dir = self.output_dir / doc_id
        save_dir.mkdir(parents=True, exist_ok=True)
        file_path = save_dir / filename

        # 渲染
        try:
            if doc_type == "docx":
                self._render_docx(template, ai_content, resolved_vars, file_path)
            elif doc_type == "xlsx":
                self._render_xlsx(template, ai_content, resolved_vars, file_path)
            elif doc_type == "pptx":
                self._render_pptx(template, ai_content, resolved_vars, file_path)
            else:
                return {"error": f"不支持的文档类型: {doc_type}", "status": "error"}
        except Exception as e:
            logger.error(f"文档渲染失败: {e}", exc_info=True)
            return {"error": f"文档渲染失败: {str(e)[:200]}", "status": "error"}

        return {
            "doc_id": doc_id,
            "template_name": template_name,
            "title": template.get("name", ""),
            "doc_type": doc_type,
            "file_path": str(file_path),
            "filename": filename,
            "status": "success",
        }

    def _scan_template_dirs(self, base_dir: Path) -> List[Path]:
        """扫描模板目录"""
        if not base_dir.exists():
            return []
        dirs = []
        for item in base_dir.iterdir():
            if item.is_dir() and (item / "template.yaml").exists():
                dirs.append(item)
        return sorted(dirs)

    def _load_template_info(self, tmpl_dir: Path, scope: str = "system") -> Optional[dict]:
        """加载模板基本信息"""
        yaml_path = tmpl_dir / "template.yaml"
        if not yaml_path.exists():
            return None

        try:
            with open(yaml_path, "r", encoding="utf-8") as f:
                config = yaml.safe_load(f)

            return {
                "name": config.get("name", tmpl_dir.name),
                "description": config.get("description", ""),
                "category": config.get("category", "其他"),
                "icon": config.get("icon", "📄"),
                "doc_type": config.get("doc_type", "docx"),
                "scope": scope,
                "variables": config.get("variables", []),
                "dir": str(tmpl_dir),
            }
        except Exception as e:
            logger.error(f"加载模板 {tmpl_dir} 失败: {e}")
            return None

    def _load_full_template(self, name: str) -> Optional[dict]:
        """加载完整模板配置"""
        for base_dir in [self.user_templates_dir, self.system_templates_dir]:
            tmpl_dir = base_dir / name
            yaml_path = tmpl_dir / "template.yaml"
            if yaml_path.exists():
                try:
                    with open(yaml_path, "r", encoding="utf-8") as f:
                        config = yaml.safe_load(f)
                    config["_dir"] = str(tmpl_dir)
                    return config
                except Exception as e:
                    logger.error(f"加载模板 {name} 失败: {e}")
        return None

    def _resolve_variables(self, template: dict, user_vars: dict) -> dict:
        """处理变量，自动填充数据库查询类型的变量"""
        resolved = {}

        for var_def in template.get("variables", []):
            var_name = var_def.get("name", "")
            var_type = var_def.get("type", "text")

            # 用户提供的值优先
            if var_name in user_vars:
                resolved[var_name] = user_vars[var_name]
                continue

            # 自动填充默认值
            if var_def.get("default"):
                resolved[var_name] = var_def["default"]
                continue

            # SQL查询类型变量，自动执行
            if var_type == "sql_query" and var_def.get("auto_fill"):
                sql = var_def.get("sql", "")
                # 替换SQL中的变量引用
                for k, v in resolved.items():
                    sql = sql.replace(f"{{{{{k}}}}}", str(v))

                try:
                    from core.database import dm_database
                    rows = dm_database.execute_query(sql)
                    resolved[var_name] = rows
                except Exception as e:
                    logger.error(f"自动查询变量 {var_name} 失败: {e}")
                    resolved[var_name] = []

        return resolved

    def _ai_fill(self, template: dict, variables: dict) -> dict:
        """AI根据Prompt填充文档内容"""
        ai_prompt = template.get("ai_prompt", "")
        doc_type = template.get("doc_type", "docx")

        # 替换Prompt中的变量引用
        for k, v in variables.items():
            placeholder = "{{" + k + "}}"
            if isinstance(v, (list, dict)):
                value_str = json.dumps(v, ensure_ascii=False, indent=2)
            else:
                value_str = str(v)
            ai_prompt = ai_prompt.replace(placeholder, value_str)

        if not ai_prompt.strip():
            ai_prompt = f"请根据以下信息生成文档内容：\n{json.dumps(variables, ensure_ascii=False)}"

        # 根据文档类型构造严格的输出格式要求
        if doc_type == "pptx":
            format_spec = """\n
严格输出如下JSON格式（不要输出任何其他文字）：
{
  "slides": [
    {"layout": "title", "title": "封面标题", "subtitle": "副标题"},
    {"layout": "content", "title": "第2页标题", "body": [{"type": "bullet", "items": ["要点1", "要点2"]}]},
    {"layout": "content", "title": "第3页标题", "body": [{"type": "text", "content": "正文内容"}]},
    {"layout": "content", "title": "第4页标题", "body": [{"type": "bullet", "items": ["要点1", "要点2", "要点3"]}]}
  ]
}

要求：
- 至少生成5-8页幻灯片
- 每页必须有title
- body中的items每页至少3个要点
- 最后1-2页必须是总结或计划"""
        elif doc_type == "xlsx":
            format_spec = """\n
严格输出如下JSON格式（不要输出任何其他文字）：
{
  "sheets": [
    {
      "name": "工作表名",
      "headers": ["列1", "列2", "列3"],
      "rows": [["数据1", "数据2", "数据3"]]
    }
  ]
}"""
        else:  # docx
            format_spec = """\n
严格输出如下JSON格式（不要输出任何其他文字）：
{
  "sections": [
    {"type": "heading", "level": 1, "text": "一、标题"},
    {"type": "paragraph", "text": "正文段落内容，至少50字"},
    {"type": "heading", "level": 2, "text": "1.1 子标题"},
    {"type": "paragraph", "text": "详细说明"},
    {"type": "bullet", "text": "要点"},
    {"type": "table", "headers": ["列1", "列2"], "rows": [["数据1", "数据2"]]}
  ]
}

要求：
- 至少6个section
- heading和paragraph交替出现
- 表格必须包含headers和rows
- 正文段落每段至少50字"""

        messages = [
            {"role": "system", "content": f"你是专业的文档生成助手。根据用户要求生成高质量的结构化文档内容。输出严格的JSON格式，不要包含任何解释文字或markdown标记。文档类型：{doc_type}"},
            {"role": "user", "content": ai_prompt + format_spec},
        ]

        try:
            result = llm_client.chat_json(messages, temperature=0.4)
            # 验证返回结构必须有对应key
            if doc_type == "pptx" and "slides" not in result and "sections" in result:
                # AI返回了sections而不是slides，自动转换
                result["slides"] = _sections_to_slides(result["sections"], template.get("name", ""))
            elif doc_type == "docx" and "sections" not in result and "slides" in result:
                # 反向转换
                result["sections"] = _slides_to_sections(result["slides"])
            return result
        except Exception as e:
            logger.error(f"AI填充失败: {e}", exc_info=True)
            # 构造最少可用内容，而不是直接报错
            if doc_type == "pptx":
                return {"slides": [
                    {"layout": "title", "title": template.get("name", "文档"), "subtitle": "AI生成失败，请重试"},
                    {"layout": "content", "title": "错误信息", "body": [{"type": "text", "content": str(e)[:200]}]},
                ]}
            else:
                return {"sections": [
                    {"type": "heading", "level": 1, "text": template.get("name", "文档")},
                    {"type": "paragraph", "text": f"AI生成失败，请重试。错误: {e}"},
                ]}

    def _render_filename(self, template: str, variables: dict, doc_type: str) -> str:
        """渲染文件名"""
        filename = template
        for k, v in variables.items():
            if isinstance(v, str):
                filename = filename.replace("{{" + k + "}}", v)

        # 内置变量
        now = datetime.now()
        filename = filename.replace("{{date}}", now.strftime("%Y%m%d"))
        filename = filename.replace("{{time}}", now.strftime("%H%M%S"))

        # 确保扩展名正确
        if not filename.endswith(f".{doc_type}"):
            filename += f".{doc_type}"

        return filename

    def _render_docx(self, template: dict, ai_content: dict,
                     variables: dict, file_path: Path):
        """渲染Word文档"""
        from renderers.docx_renderer import render_docx

        # 检查是否有模板文件
        tmpl_dir = template.get("_dir", "")
        template_file = Path(tmpl_dir) / template.get("template_file", "") if tmpl_dir else None

        if template_file and template_file.exists():
            # 从模板文件渲染
            self._render_from_template_docx(template_file, ai_content, variables, file_path)
        else:
            # 从AI内容直接生成
            doc_content = {
                "title": template.get("name", "文档"),
                "sections": ai_content.get("sections", []),
            }
            render_docx(doc_content, file_path)

    def _render_from_template_docx(self, template_file: Path, ai_content: dict,
                                    variables: dict, file_path: Path):
        """从Word模板文件渲染（替换占位符）"""
        from docx import Document

        doc = Document(str(template_file))

        # 替换段落中的占位符
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                text = run.text
                modified = False

                # 替换变量占位符 {{var}}
                for k, v in variables.items():
                    placeholder = "{{" + k + "}}"
                    if placeholder in text:
                        if isinstance(v, (list, dict)):
                            text = text.replace(placeholder, json.dumps(v, ensure_ascii=False)[:100])
                        else:
                            text = text.replace(placeholder, str(v))
                        modified = True

                # 替换AI内容占位符 {{ai_section_N}}
                ai_sections = ai_content.get("sections", [])
                for i, section in enumerate(ai_sections, 1):
                    placeholder = f"{{{{ai_section_{i}}}}}"
                    if placeholder in text:
                        text = text.replace(placeholder, section.get("text", ""))
                        modified = True

                if modified:
                    run.text = text

        # 替换表格中的占位符
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            text = run.text
                            for k, v in variables.items():
                                placeholder = "{{" + k + "}}"
                                if placeholder in text:
                                    text = text.replace(placeholder, str(v) if not isinstance(v, (list, dict)) else "")
                            run.text = text

        doc.save(str(file_path))

    def _render_xlsx(self, template: dict, ai_content: dict,
                     variables: dict, file_path: Path):
        """渲染Excel文档"""
        from renderers.xlsx_renderer import render_xlsx
        doc_content = {
            "title": template.get("name", "表格"),
            "sheets": ai_content.get("sheets", []),
        }
        render_xlsx(doc_content, file_path)

    def _render_pptx(self, template: dict, ai_content: dict,
                     variables: dict, file_path: Path):
        """渲染PPT文档"""
        from renderers.pptx_renderer import render_pptx
        doc_content = {
            "title": template.get("name", "演示文稿"),
            "slides": ai_content.get("slides", []),
        }
        render_pptx(doc_content, file_path)


# 全局单例
template_engine = TemplateEngine()
