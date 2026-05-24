from __future__ import annotations
"""文件解析模块 - 从上传文件中提取文本内容作为对话上下文

支持格式：
- Word (.docx)
- Excel (.xlsx)
- PowerPoint (.pptx)
- PDF (.pdf)
- 图片 (.png/.jpg/.jpeg/.gif/.bmp/.webp) — 使用OCR或视觉模型
- 纯文本 (.txt/.md/.csv/.json/.yaml/.yml)
"""
import base64
import logging
import os
from pathlib import Path
from typing import Optional

from config.settings import settings

logger = logging.getLogger(__name__)

# 支持的文件扩展名
SUPPORTED_EXTENSIONS = {
    # 文档
    ".docx": "word",
    ".xlsx": "excel",
    ".pptx": "ppt",
    ".pdf": "pdf",
    # 图片
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".bmp": "image",
    ".webp": "image",
    # 纯文本
    ".txt": "text",
    ".md": "text",
    ".csv": "text",
    ".json": "text",
    ".yaml": "text",
    ".yml": "text",
}


def get_file_type(filename: str) -> Optional[str]:
    """根据文件名获取文件类型"""
    ext = Path(filename).suffix.lower()
    return SUPPORTED_EXTENSIONS.get(ext)


def is_supported(filename: str) -> bool:
    """检查文件是否支持"""
    return get_file_type(filename) is not None


def extract_text(file_path: str, filename: str = None) -> str:
    """从文件中提取文本内容

    Args:
        file_path: 文件路径
        filename: 原始文件名（用于判断类型）

    Returns:
        提取的文本内容
    """
    fname = filename or os.path.basename(file_path)
    file_type = get_file_type(fname)

    if file_type is None:
        # 尝试作为纯文本
        return _extract_plain_text(file_path)

    extractors = {
        "word": _extract_docx,
        "excel": _extract_xlsx,
        "ppt": _extract_pptx,
        "pdf": _extract_pdf,
        "image": _extract_image,
        "text": _extract_plain_text,
    }

    extractor = extractors.get(file_type)
    if extractor:
        return extractor(file_path)

    return _extract_plain_text(file_path)


def extract_image_base64(file_path: str) -> Optional[str]:
    """提取图片的base64编码（用于视觉模型）"""
    try:
        with open(file_path, "rb") as f:
            return base64.b64encode(f.read()).decode("utf-8")
    except Exception as e:
        logger.error(f"读取图片失败: {e}")
        return None


def get_image_media_type(filename: str) -> str:
    """获取图片的MIME类型"""
    ext = Path(filename).suffix.lower()
    types = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
        ".webp": "image/webp",
    }
    return types.get(ext, "image/png")


# ===== 具体解析器 =====

def _extract_docx(file_path: str) -> str:
    """提取Word文档文本"""
    from docx import Document
    doc = Document(file_path)

    parts = []
    # 段落
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # 表格
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n\n".join(parts)


def _extract_xlsx(file_path: str) -> str:
    """提取Excel文本"""
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True)

    parts = []
    for ws in wb.worksheets:
        parts.append(f"=== 工作表: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))

    wb.close()
    return "\n".join(parts)


def _extract_pptx(file_path: str) -> str:
    """提取PPT文本"""
    from pptx import Presentation
    prs = Presentation(file_path)

    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    slide_parts.append(" | ".join(cells))
        if slide_parts:
            parts.append(f"=== 第{i}页 ===\n" + "\n".join(slide_parts))

    return "\n\n".join(parts)


def _extract_pdf(file_path: str) -> str:
    """提取PDF文本"""
    # 方式1: PyMuPDF
    try:
        import fitz
        doc = fitz.open(file_path)
        text = "\n\n".join(page.get_text() for page in doc)
        doc.close()
        if text.strip():
            return text
    except ImportError:
        pass

    # 方式2: pdfplumber
    try:
        import pdfplumber
        texts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    texts.append(t)
        if texts:
            return "\n\n".join(texts)
    except ImportError:
        pass

    # 方式3: PyPDF2
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(file_path)
        texts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                texts.append(t)
        if texts:
            return "\n\n".join(texts)
    except ImportError:
        pass

    return "[PDF解析失败：缺少依赖库，请安装 PyMuPDF/pdfplumber/PyPDF2 之一]"


def _extract_image(file_path: str) -> str:
    """图片文件 - 返回标记，由对话逻辑用视觉模型处理"""
    filename = os.path.basename(file_path)
    return f"[图片文件: {filename}]"


def _extract_plain_text(file_path: str) -> str:
    """纯文本提取"""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        try:
            with open(file_path, "r", encoding="gbk", errors="ignore") as f:
                return f.read()
        except Exception as e:
            return f"[文件读取失败: {e}]"
